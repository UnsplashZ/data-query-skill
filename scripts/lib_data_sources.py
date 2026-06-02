#!/usr/bin/env python3
"""Shared helpers for internal-data-query scripts."""

from __future__ import annotations

import csv
import json
import os
import re
import sys
import time
import urllib.error
import urllib.parse
import urllib.request
from pathlib import Path
from typing import Any

SUPPORTED_ENGINES = ("clickhouse", "odps", "mysql", "metabase")

REQUIRED_PROFILE_FIELDS: dict[str, tuple[str, ...]] = {
    "clickhouse": ("host", "database", "username", "password"),
    "odps": ("endpoint", "project", "access_id", "access_key"),
    "mysql": ("host", "database", "username", "password"),
    "metabase": ("base_url",),
}

PLACEHOLDER_MARKERS = (
    "example.",
    ".example",
    "your_",
    "placeholder",
    "changeme",
    "change_me",
    "readonly_user",
    "${",
)


def parse_env_file(path: Path | None) -> dict[str, str]:
    values: dict[str, str] = {}
    if not path:
        return values
    if not path.exists():
        raise FileNotFoundError(f"env file not found: {path}")
    for raw in path.read_text(encoding="utf-8").splitlines():
        line = raw.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        values[key.strip()] = value.strip().strip('"').strip("'")
    return values


def _expand_env(value: Any, env: dict[str, str]) -> Any:
    if isinstance(value, str):
        pattern = re.compile(r"\$\{([A-Za-z_][A-Za-z0-9_]*)\}")
        return pattern.sub(lambda m: env.get(m.group(1), os.environ.get(m.group(1), "")), value)
    if isinstance(value, dict):
        return {k: _expand_env(v, env) for k, v in value.items()}
    if isinstance(value, list):
        return [_expand_env(v, env) for v in value]
    return value


def default_config_candidates() -> list[Path]:
    candidates: list[Path] = []
    env_path = os.environ.get("INTERNAL_DATA_QUERY_CONFIG")
    if env_path:
        candidates.append(Path(env_path).expanduser())
    candidates.extend(
        [
            Path.home() / ".internal-data-query" / "data-sources.yaml",
            Path("local/data-sources.yaml"),
            Path("data-query-work/config/data-sources.yaml"),
        ]
    )
    return candidates


def resolve_config_path(path: Path | None) -> Path | None:
    if path:
        return path.expanduser()
    for candidate in default_config_candidates():
        if candidate.exists():
            return candidate
    return None


def load_yaml_config(path: Path | None, env: dict[str, str]) -> dict[str, Any]:
    path = resolve_config_path(path)
    if not path:
        return {}
    if not path.exists():
        raise FileNotFoundError(f"config file not found: {path}")
    try:
        import yaml  # type: ignore
    except ImportError as exc:
        raise RuntimeError(
            "PyYAML is required for --config YAML files. Install dependencies with: "
            "python -m pip install -r requirements.txt"
        ) from exc
    data = yaml.safe_load(path.read_text(encoding="utf-8")) or {}
    return _expand_env(data, env)


def load_config_with_path(config_path: Path | None, env_file: Path | None = None) -> tuple[Path | None, dict[str, Any]]:
    env = {**os.environ, **parse_env_file(env_file)}
    resolved = resolve_config_path(config_path)
    return resolved, load_yaml_config(resolved, env)


def iter_profile_configs(config: dict[str, Any], engine: str | None = None) -> list[tuple[str, str, dict[str, Any]]]:
    rows: list[tuple[str, str, dict[str, Any]]] = []
    profiles = config.get("profiles") or {}
    engines = [engine] if engine else list(SUPPORTED_ENGINES)
    for item in engines:
        for profile, cfg in (profiles.get(item) or {}).items():
            rows.append((item, str(profile), dict(cfg or {})))
    return rows


def is_placeholder_value(value: Any) -> bool:
    if value in (None, ""):
        return False
    text = str(value).strip().lower()
    return any(marker in text for marker in PLACEHOLDER_MARKERS)


def profile_missing_fields(engine: str, cfg: dict[str, Any]) -> list[str]:
    missing = [field for field in REQUIRED_PROFILE_FIELDS.get(engine, ()) if cfg.get(field) in (None, "")]
    if engine == "metabase" and not (cfg.get("api_key") or cfg.get("session_id") or (cfg.get("username") and cfg.get("password"))):
        missing.append("api_key|session_id|username+password")
    return missing


def profile_placeholder_fields(engine: str, cfg: dict[str, Any]) -> list[str]:
    fields = set(REQUIRED_PROFILE_FIELDS.get(engine, ()))
    if engine == "metabase":
        fields.update({"api_key", "session_id", "username", "password"})
    return sorted(field for field in fields if is_placeholder_value(cfg.get(field)))


def bool_value(value: Any, default: bool = False) -> bool:
    if value is None:
        return default
    if isinstance(value, bool):
        return value
    return str(value).strip().lower() in {"1", "true", "yes", "y", "on"}


def load_profile(engine: str, profile: str, config_path: Path | None, env_file: Path | None) -> dict[str, Any]:
    env = {**os.environ, **parse_env_file(env_file)}
    resolved_config_path = resolve_config_path(config_path)
    config = load_yaml_config(resolved_config_path, env)
    profiles = (config.get("profiles") or {}).get(engine) or {}
    cfg = dict(profiles.get(profile) or {})

    prefix = engine.upper()
    aliases = {
        "clickhouse": {
            "username": ["CLICKHOUSE_USERNAME", "CLICKHOUSE_USER"],
            "password": ["CLICKHOUSE_PASSWORD"],
            "host": ["CLICKHOUSE_HOST"],
            "port": ["CLICKHOUSE_PORT"],
            "database": ["CLICKHOUSE_DATABASE"],
            "secure": ["CLICKHOUSE_SECURE"],
        },
        "mysql": {
            "username": ["MYSQL_USERNAME", "MYSQL_USER"],
            "password": ["MYSQL_PASSWORD"],
            "host": ["MYSQL_HOST"],
            "port": ["MYSQL_PORT"],
            "database": ["MYSQL_DATABASE"],
            "charset": ["MYSQL_CHARSET"],
        },
        "metabase": {
            "base_url": ["METABASE_BASE_URL"],
            "api_key": ["METABASE_API_KEY"],
            "session_id": ["METABASE_SESSION_ID"],
            "username": ["METABASE_USERNAME"],
            "password": ["METABASE_PASSWORD"],
        },
        "odps": {
            "endpoint": ["ODPS_ENDPOINT"],
            "project": ["ODPS_PROJECT"],
            "access_id": ["ODPS_ACCESS_ID"],
            "access_key": ["ODPS_ACCESS_KEY"],
            "tunnel_endpoint": ["ODPS_TUNNEL_ENDPOINT"],
        },
    }
    for key, names in aliases.get(engine, {}).items():
        if cfg.get(key) in (None, ""):
            for name in names:
                if env.get(name):
                    cfg[key] = env[name]
                    break
    if not cfg:
        raise RuntimeError(
            f"No config found for engine={engine} profile={profile}.\n"
            "Run: python scripts/setup_connections.py\n"
            "Then retry, or pass --config ~/.internal-data-query/data-sources.yaml.\n"
            "Scripts also read $INTERNAL_DATA_QUERY_CONFIG, local/data-sources.yaml, engine .env files, "
            "or process environment variables."
        )
    return cfg


def read_sql(sql: str | None, sql_file: Path | None) -> str:
    if sql:
        return sql
    if sql_file:
        return sql_file.read_text(encoding="utf-8")
    data = sys.stdin.read()
    if data.strip():
        return data
    raise RuntimeError("Provide --sql, --sql-file, or SQL on stdin.")


def ensure_output_dir(path: Path) -> Path:
    path.mkdir(parents=True, exist_ok=True)
    return path


def write_csv(path: Path, columns: list[str], rows: list[dict[str, Any]]) -> None:
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=columns, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def write_xlsx(path: Path, columns: list[str], rows: list[dict[str, Any]]) -> None:
    try:
        from openpyxl import Workbook  # type: ignore
    except ImportError as exc:
        raise RuntimeError("openpyxl is required for xlsx output. Use --output-format csv or install openpyxl.") from exc
    wb = Workbook()
    ws = wb.active
    ws.title = "result"
    ws.append(columns)
    for row in rows:
        ws.append([row.get(col) for col in columns])
    wb.save(path)


def export_rows(output_dir: Path, stem: str, fmt: str, rows: list[dict[str, Any]], columns: list[str] | None = None) -> Path:
    ensure_output_dir(output_dir)
    if columns is None:
        seen: list[str] = []
        for row in rows:
            for key in row.keys():
                if key not in seen:
                    seen.append(key)
        columns = seen
    path = output_dir / f"{stem}.{fmt}"
    if fmt == "csv":
        write_csv(path, columns, rows)
    elif fmt == "xlsx":
        write_xlsx(path, columns, rows)
    else:
        raise ValueError(f"unsupported output format: {fmt}")
    return path


def now_stem(prefix: str) -> str:
    return f"{time.strftime('%Y%m%d-%H%M%S')}-{re.sub(r'[^A-Za-z0-9_-]+', '-', prefix).strip('-') or 'query'}"


def metabase_headers(cfg: dict[str, Any]) -> dict[str, str]:
    headers = {"Content-Type": "application/json"}
    if cfg.get("api_key"):
        headers["x-api-key"] = str(cfg["api_key"])
    elif cfg.get("session_id"):
        headers["X-Metabase-Session"] = str(cfg["session_id"])
    return headers


def metabase_request(cfg: dict[str, Any], method: str, path: str, payload: dict[str, Any] | None = None) -> Any:
    base_url = str(cfg.get("base_url") or "").rstrip("/")
    if not base_url:
        raise RuntimeError("Metabase base_url is required.")
    placeholders = profile_placeholder_fields("metabase", cfg)
    if placeholders:
        raise RuntimeError(
            "Metabase profile still contains placeholder values: "
            + ", ".join(placeholders)
            + ". Placeholder config is only for setup parsing/discovery smoke. "
            "Use --mock-file for offline Metabase checks, or replace the local config with a real readonly Metabase URL and credentials."
        )
    if not (cfg.get("api_key") or cfg.get("session_id") or (cfg.get("username") and cfg.get("password"))):
        raise RuntimeError(
            "Metabase credentials are required for real API calls. Configure api_key, session_id, or username+password in the local readonly profile; "
            "for installation smoke use check_connections/discover/help commands or Metabase --mock-file fixtures."
        )
    if not (cfg.get("api_key") or cfg.get("session_id")) and cfg.get("username") and cfg.get("password"):
        session = metabase_login(cfg)
        cfg["session_id"] = session
    data = json.dumps(payload or {}).encode("utf-8") if payload is not None else None
    req = urllib.request.Request(base_url + path, data=data, headers=metabase_headers(cfg), method=method)
    try:
        with urllib.request.urlopen(req, timeout=120) as resp:
            body = resp.read().decode("utf-8")
            return json.loads(body) if body else None
    except urllib.error.HTTPError as exc:
        detail = exc.read().decode("utf-8", errors="ignore")
        raise RuntimeError(
            f"Metabase request failed: base_url={base_url}, http_status={exc.code}.\n"
            f"Detail: {detail[:500]}\n"
            "Check base_url, credentials/session/API key, object id, permissions, and network/VPN access."
        ) from exc
    except urllib.error.URLError as exc:
        raise RuntimeError(
            f"Metabase connection failed: base_url={base_url}, error={exc.reason}.\n"
            "Check base_url, DNS/network/VPN, proxy settings, and whether Metabase is reachable from this machine."
        ) from exc


def metabase_login(cfg: dict[str, Any]) -> str:
    base_url = str(cfg.get("base_url") or "").rstrip("/")
    payload = {"username": cfg.get("username"), "password": cfg.get("password")}
    data = json.dumps(payload).encode("utf-8")
    req = urllib.request.Request(
        base_url + "/api/session",
        data=data,
        headers={"Content-Type": "application/json"},
        method="POST",
    )
    try:
        with urllib.request.urlopen(req, timeout=60) as resp:
            body = json.loads(resp.read().decode("utf-8"))
    except urllib.error.HTTPError as exc:
        detail = exc.read().decode("utf-8", errors="ignore")
        raise RuntimeError(
            f"Metabase login failed: base_url={base_url}, http_status={exc.code}.\n"
            f"Detail: {detail[:500]}\n"
            "Check username/password, API key preference, and Metabase authentication settings."
        ) from exc
    except urllib.error.URLError as exc:
        raise RuntimeError(
            f"Metabase login connection failed: base_url={base_url}, error={exc.reason}.\n"
            "Check base_url, DNS/network/VPN, proxy settings, and whether Metabase is reachable."
        ) from exc
    return body["id"]


def exit_with_error(exc: BaseException) -> int:
    print(f"ERROR: {exc}", file=sys.stderr)
    return 1


def metabase_result_to_rows(result: Any) -> tuple[list[str], list[dict[str, Any]]]:
    if isinstance(result, list):
        rows = [dict(item) for item in result]
        columns = list(rows[0].keys()) if rows else []
        return columns, rows
    data = (result or {}).get("data") or {}
    cols = [col.get("name") or col.get("display_name") or f"col_{i+1}" for i, col in enumerate(data.get("cols") or [])]
    rows = [dict(zip(cols, row)) for row in data.get("rows") or []]
    return cols, rows
